# -*- coding: utf-8 -*-
"""
Bilingual conversion for the pest-control exam bank.
All English text is authored by the LLM (no external/online translation API).
Input : questions_backup_20260929.json
Output: questions_converted_draft.json  (UTF-8, no BOM, same structure)
        ../outputs/question_bank_final_preview.docx
"""
import json, re, os, copy

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "questions_backup_20260929.json")
OUT_JSON = os.path.join(BASE, "questions_converted_draft.json")
OUT_DOCX = os.path.join(BASE, "..", "outputs", "question_bank_final_preview.docx")

data = json.load(open(SRC, encoding="utf-8"))

COMBO_RE = re.compile(r"[IVX]+\s*\.?\s*\+")

def is_combo_options(opts):
    return any(COMBO_RE.search(o) for o in opts)

def has_chinese(s):
    return any("\u4e00" <= c <= "\u9fff" for c in s)

STMT_RE = re.compile(r"([IVX]+)\s*\.")

def parse_statements(qtext):
    matches = list(STMT_RE.finditer(qtext))
    if not matches:
        return qtext.strip(), []
    stem = qtext[: matches[0].start()].strip()
    stmts = []
    for i, m in enumerate(matches):
        roman = m.group(1)
        start = m.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(qtext)
        txt = qtext[start:end].strip()
        stmts.append((roman, txt))
    return stem, stmts

def max_roman_in_options(opts):
    mx = 0
    order = {"I":1,"II":2,"III":3,"IV":4,"V":5,"VI":6,"VII":7,"VIII":8}
    for o in opts:
        for m in re.findall(r"([IVX]+)", o):
            if m in order:
                mx = max(mx, order[m])
    return mx

# ----------------------------------------------------------------------------
# Translation tables (LLM-authored).
#   statement question -> 's' (stem english) + 'st' {roman: english}
#   plain question     -> 'q' (full english)   + 'o'  [per-option english or None]
#   plain-option statement question -> also 'o' for the options.
#   'o' = None  => option kept as-is (pure number/English, or already bilingual).
# ----------------------------------------------------------------------------
T = {
"1": {
 0: {"q":"According to Hong Kong's Agriculture, Fisheries and Conservation Department (AFCD), which Ordinance regulates the use of pesticides (insecticides)?",
     "o":["Pesticides Ordinance (Cap. 133)","Pesticides Ordinance (Cap. 144)","Pesticides Ordinance (Cap. 155)","Pesticides Ordinance (Cap. 166)"]},
 1: {"s":"What are the consequences of excessive use of insecticides?",
     "st":{"I":"Adverse effects on beneficial insects","II":"Long-term health effects",
           "III":"Insects develop resistance to insecticides","IV":"Residues left in food"}},
 2: {"q":"What is the English abbreviation of the government department, the Agriculture, Fisheries and Conservation Department, that regulates insecticides?",
     "o":[None,None,None,None]},
 3: {"q":"According to AFCD regulations, what must a registered insecticide display on the label on its container?",
     "o":["Registered trademark","Registration number","Drug efficacy","All of the above"]},
 4: {"q":"All insecticides have a ________ (MSDS), which contains information on the toxicity of the insecticide and its health effects.",
     "o":["Pharmaceutical safety data sheet","Material Safety Data Sheet (MSDS)","Material and pharmaceutical health data sheet","Insecticide content data sheet"]},
 5: {"q":"What is the correct order of the 6 steps of Integrated Pest Management (IPM)?",
     "o":["Monitoring, identification, inspection, review, education, implementation",
          "Identification, inspection, implementation, monitoring, review, education",
          "Education, implementation, identification, monitoring, inspection, review",
          "Inspection, identification, monitoring, implementation, review, education"]},
 6: {"q":"Must all insecticides in Hong Kong be registered before they can be legally used?",
     "o":["Yes","No","Not necessarily","Depends on the intended use"]},
 7: {"q":"What is Integrated Pest Management (IPM)?\nA continuous approach that combines ____:____:____:____ tools to manage pests while reducing risks to the economy, human health and the environment.",
     "o":["Environmental, ecological, organic, improved living","Biological, natural enemies, environmental improvement, safe application",
          "Physiological, biological, physical, chemical","Biological, cultural, physical, chemical"]},
 8: {"q":"Where can the registration number of an insecticide be found?",
     "o":["Front of the label sticker on the container","Back of the label sticker","Bottom of the container","Top of the container"]},
 9: {"q":"Bed bug control is often affected by _____, due to widespread and excessive use of ____?",
     "o":["Environment, insecticide","Client, data","Resistance, chlordane (variant)","Resistance, DDT"]},
 10: {"q":"Under Integrated Pest Management, how do we inspect for pests?",
     "o":["Identify pests and their conditions, and signs of their presence","Trust the client's intuition",
          "Pay attention to the client's reaction","Spray insecticide directly"]},
 11: {"q":"Which of the following is NOT a type of pesticide?",
     "o":["Fungicide","Acaricide (miticide)","Herbicide","Deodorant"]},
 12: {"q":"According to IPM, which method is NOT included in implementation?",
     "o":["Chemical method","Biological method","Physical method","Psychological method"]},
 13: {"q":"Installing traps in suspected infested areas and conducting regular inspections \u2014 under IPM, what is this action?",
     "o":["Implementation","Monitoring","Review","Identification"]},
 14: {"s":"Which of the following is the correct method to educate clients to prevent mosquitoes from returning?",
     "st":{"I":"Install mosquito nets","II":"Maintain a clean and hygienic environment",
           "III":"Clear stagnant water","IV":"Open windows to maintain ventilation"}},
 15: {"q":"According to what criteria are insecticides diluted?",
     "o":["According to the client's requests and instructions","According to the dosage stated on the registered label",
          "According to the label with appropriate adjustment","According to one's own experience and site conditions"]},
 16: {"q":"Under Hong Kong law, which department and which Ordinance regulate the use of pesticides (insecticides)?",
     "o":["AFCD and the Pesticides Ordinance (Cap. 133)","AFCD and the Pesticides Ordinance (Cap. 133)",
          "Agriculture, Fisheries and Conservation Department and the Pesticides Ordinance (Cap. 133)",
          "Director of AFCD and the Pesticides Ordinance (Cap. 133)"]},
 17: {"q":"Which of the following is NOT information that an insecticide label tells us?",
     "o":["Source of origin","Production process","Where and how to apply it","Storage location"]},
 18: {"q":"Using mosquito sand (larvicide) to kill mosquito larvae is which IPM implementation method?",
     "o":["Physiological","Physical","Chemical","Psychological"]},
 19: {"q":"Why does the company advise clients not to seal drain openings with tape to prevent American cockroaches from entering?",
     "o":["Not aesthetically pleasing","Affects the ecology","Blocks drainage","Business consideration"]},
 20: {"q":"By definition, the term 'pesticide' covers which of the following target organisms?",
     "o":["Only insects","Only insects and rodents",
          "Any harmful organism including insects, rodents, birds, nematodes, bacteria, fungi, weeds, etc.","Only weeds and fungi"]},
 21: {"q":"Which of the following is a consequence of 'excessive use of insecticides'?",
     "o":["Pests are eliminated faster and never return","Resistance","Increase in bee populations","Environmental pollution"]},
 22: {"q":"Why is Imidacloprid criticised globally \u2014 because it harms which group of organisms?",
     "o":["Rodents","Bee populations","Birds","Koalas"]},
 23: {"q":"In 'physical methods', examples of 'habit modification' include?",
     "o":["Cleaning and removing pest habitats","Spraying large amounts of chemicals","Releasing natural enemies","Placing bait"]},
 24: {"q":"Regarding sticky traps and pheromones, which is correct?",
     "o":["Sticky traps are a good non-chemical control; adding pheromone granules enhances efficacy",
          "Pheromones attract all pests","Sticky traps can only be used outdoors","Pheromones do not help traps"]},
 25: {"q":"When using chemical methods (insecticides), what principle should be followed?",
     "o":["Spray everywhere as much as possible","Targeted, controlled use","The higher the concentration the better","Follow superiors' instructions"]},
 26: {"q":"In 'biological methods', synthetic insect growth regulators (IGR) are mainly used to control which pest?",
     "o":["Rodents","Termites","Birds","Weeds"]},
 27: {"q":"The ichneumon wasp (corn wasp) is a natural enemy of cockroaches; how does it kill them?",
     "o":["Its acidic odour can effectively kill cockroaches",
          "It naturally eats cockroaches, so where it exists cockroaches cannot survive",
          "It lays eggs inside cockroach eggs; the larvae hatch and eat the cockroach larvae",
          "It continuously emits ultra-low sound waves that affect cockroaches mentally"]},
},
"2": {
 0: None, 1: None, 2: None, 3: None, 4: None,
 5: {"q":"What mode of action does BioKill Universal Insecticide (BKU) use to kill pests?",
     "o":[None,None,None,None]},
 6: {"q":"What is the main difference between BioKill Universal Insecticide and BioKill Long Lasting Extra Insecticide?",
     "o":["The microcapsule of Long Lasting Extra periodically releases active insecticide ingredients",
          "Only Universal can be widely applied to residences, offices, factories, restaurants, hospitals, zoos, plants",
          "Long Lasting Extra can be applied to articles and surroundings of warm-blooded animals",
          "None; both can be used in Hand pump, Turbo and Backpack"]},
 7: {"q":"How does Temephos effectively stop mosquito growth in treatment?",
     "o":["Effectively kills adult mosquitoes","Effectively absorbs oxygen from water",
          "Effectively stops mosquito larvae growth, interrupting the life cycle","Effectively kills all flying insects"]},
 8: {"q":"According to Fendona's label guidance, where should the main application be?",
     "o":["Drains, cracks, persistently damp (deteriorated) locations","Sewers, slopes, hidden places",
          "Broad-area spraying everywhere","Damp, pest-harassed, dirty places and indoor locations"]},
 9: {"q":"Why should Premise not be sprayed directly on plant surfaces?",
     "o":["It damages plant roots after penetrating the soil","Affects plant nutrient absorption",
          "It causes great harm to bees","Prevents plants from photosynthesis"]},
 10: {"q":"What type of cockroach is the so-called 'cockroach sandwich' (roach bait gel) effective against?",
     "o":["American cockroach","Brown-banded cockroach","Asian cockroach","German cockroach"]},
 11: {"q":"If a home has pests such as cockroaches, ants, booklice and silverfish, with no pets, which product would we suggest the client buy?",
     "o":["BioKill Mite Spray BKA","BioKill Universal Insecticide BKU","BioKill Long Lasting Extra BKL","BioKill Safe Disinfectant BKS"]},
 12: {"q":"According to the LD50 theory, which of the following numbers represents the highest toxicity?",
     "o":[None,None,None,None]},
 13: {"q":"When rodent bait stations are used for external perimeter control, approximately how many metres apart are they generally installed around a building?",
     "o":["5-10 m","10-15 m","15-20 m","20-25 m"]},
 14: {"q":"Fumigation poisoning uses the ______ of the active ingredient to enter the pest's body through the respiratory system (spiracles or body wall).",
     "o":["Gas","Solid","Foam","Liquid"]},
 15: {"s":"What should we pay attention to when using Phantom to treat bed bugs?",
     "st":{"I":"Must not be sprayed directly on occupant-contact surfaces",
           "II":"Broad-area spraying everywhere",
           "III":"Spray on mattress edges and bed frame",
           "IV":"Spray on skirting boards, gaps and cracks"}},
 16: {"q":"The oral toxicity LD50 of BKU in rat testing is",
     "o":[None,None,None,None]},
 17: {"q":"The exposure period for pesticide 'chronic toxicity' is usually?",
     "o":["Within 24 hours","1-2 years","1-2 weeks","1-2 months"]},
 18: {"q":"Percutaneous penetration through the skin is the most common poisoning route. Which body part absorbs the chemical fastest?",
     "o":["Forearm","Head and hair","Scrotum","Palm"]},
 19: {"q":"The active ingredient of BioKill Universal Insecticide is",
     "o":[None,None,None,None]},
 20: {"q":"The ingredients of BioKill Long Lasting Extra Insecticide include?",
     "o":[None,None,None,None]},
 21: {"q":"Which year did the World Health Organization (WHO) recommend the pesticide toxicity classification?",
     "o":[None,None,None,None]},
 22: {"q":"What does LD50 mean?",
     "o":["The amount of chemical that kills 100% of animals","The concentration that makes half ill",
          "Maximum safe dose","The amount that kills 50% of a given animal species"]},
 23: {"q":"Which method is usually used to express pesticide 'acute toxicity'?",
     "o":["LD50 or LC50","pH value","Concentration percentage","Shelf-life length"]},
},
"3": {
 0: {"q":"What is the main activity range of drywood termites that we normally treat? Choose the correct answer.",
     "o":["Dry forests, jungles, semi-finished wood products","Trees, shrubs, wood sculptures, wooden door frames and doors",
          "Damp wood, mouldy trees, around the house exterior","Tree roots, wood sculptures, anywhere indoors"]},
 1: {"q":"What diameter drill bit do we select for termite treatment drilling?",
     "o":[None,None,None,None]},
 2: {"q":"Why is use of insecticide generally not recommended for clients during termite treatment?",
     "o":["It is purely a commercial decision to charge higher fees",
          "Ordinary household insecticides have a strong odour that causes discomfort; BioKill products may be used",
          "Household insecticides generally contain coagulants that aggravate termite attack and damage",
          "Household insecticides generally contain repellents that aggravate termite spread"]},
 3: {"q":"What is the most suitable temperature for termite survival, in degrees Celsius?",
     "o":["30 to 40 degrees Celsius","Minus 3 to 37 degrees Celsius","24 to 35 degrees Celsius","20 to 30 degrees Celsius"]},
 4: {"q":"The king termite plays a very important role in a colony; in the nest he helps build a chamber for mating with the queen, and in the early formation of the colony assists with what? Choose the correct answer.",
     "o":["Attack intruders","Feed soldier termites","Feed nymphs","Feed the queen"]},
 5: {"q":"Which describes the characteristics of the termite queen?",
     "o":["About 15 mm long; lays about 100 eggs per day","About 25 mm long; lays about 1000 eggs per day",
          "About 35 mm long; lays about 1000 eggs per day","About 45 mm long; lays about 1500 eggs per day"]},
 6: {"q":"What is the shape of drywood termite faecal pellets?",
     "o":["Hexagonal","Triangular","Cylindrical","Square"]},
 7: {"q":"What is the job of soldier termites in a Formosan termite nest?",
     "o":["Attack intruders","Feed nymphs (BB)","Care for the queen","Distribute food"]},
 8: {"s":"What are the benefits of using the Sentricon colony elimination system?",
     "st":{"I":"Provides long-term protection","II":"Eliminates the entire termite colony",
           "III":"Reduces treatment risk","IV":"No repeated chemical spraying"}},
 9: {"q":"Termites belong to ___ animals; they have connected exoskeleton and cuticle?",
     "o":["Arthropod","Jointed-limb","Arthropod (articulated)","Disjointed"]},
 10: {"q":"What is the job or characteristic of worker termites in a Formosan termite nest?",
     "o":["Build and repair the nest and tunnels","They are responsible for caring for and feeding different members",
          "Worker termites are about 5 mm long","All of the above"]},
 11: {"q":"We mainly target Formosan termites / Taiwan subterranean termites; globally how ___ aggressive is this species?",
     "o":["First","Second","Third","Fourth"]},
 12: {"q":"Termite liquid chemical forms a chemical barrier via injection to protect wooden structures and control invasion; which of the following cannot use this method?",
     "o":["Active termite range","All wood products","All electrified objects","All fixed cabinets (built-in) and sinks"]},
 13: {"s":"How can we know a termite nest has been eliminated by Sentricon?",
     "st":{"I":"No more worker termites","II":"Bait consumption reaches 100%",
           "III":"Bait consumption is less than before","IV":"No feeding for a period of time"}},
 14: {"q":"Sentricon controls termites via chitin synthesis inhibition, causing moulting failure and eliminating the whole colony. Which caste does it mainly target?",
     "o":["Queen","King","Soldier","Worker"]},
 15: {"q":"Which is the most active season for termite colonies in Hong Kong?",
     "o":["Autumn, Winter","Summer, Winter","Spring, Summer","Active at any time"]},
 16: {"q":"During residential termite liquid injection, at what distance apart do we drill holes?",
     "o":["1 to 2 metres","3 to 4 metres","5 to 6 metres","7 to 8 metres"]},
 17: {"s":"What is the function of the termite cuticle?",
     "st":{"I":"Protect the body","II":"Support","III":"Prevent bacterial invasion","IV":"Prevent body water loss"}},
 18: {"q":"How far from the nest can worker termites travel to forage?",
     "o":["100 feet","300 feet","500 feet","800 feet"]},
 19: {"q":"Which is the most suitable temperature for termite survival?",
     "o":["15 to 40 degrees Celsius","Minus 3 to 37 degrees Celsius","24 to 35 degrees Celsius","20 to 30 degrees Celsius"]},
 20: {"q":"How long can a termite queen live?",
     "o":["5 to 8 years","8 to 12 years","8 to 15 years","10 to 15 years"]},
 21: {"q":"Which is NOT the work of worker termites?",
     "o":["Build and repair the nest and tunnels","Protect the nest and attack intruders",
          "Forage","Care for and feed different members"]},
 22: {"q":"During inspection, when soldier termites are disturbed, what behaviour do they show?",
     "o":["Notify the queen and immediately leave the nest","Move all larvae to a safe place",
          "Make a sound with their mandibles as an alarm","Stop worker activity and attack intruders together"]},
 23: {"s":"What are the characteristics of reproductive termites?",
     "st":{"I":"Good at flying","II":"Mate","III":"Called flying ants",
           "IV":"Appear when the colony develops to 4-5 years","V":"Die within a few hours"}},
 24: {"q":"Which of the following descriptions of 'wetwood termites' and 'drywood termites' is correct?",
     "o":["Causes the most severe structural damage","Mainly nest inside wood; colony significantly smaller than soil-nesting termites",
          "Must contact soil to survive","Faecal pellets are round"]},
 25: {"q":"What is the internal nickname of 'Yellow-winged termite' (Macrotermes Barneyi)?",
     "o":["Formosan termite","Soil termite","Subterranean termite","Winged termite"]},
 26: {"q":"What are the approximate lower and upper temperature limits for termite survival?",
     "o":["-3 C to 37 C","0 C to 30 C","5 C to 40 C","-10 C to 25 C"]},
 27: {"q":"Formosan termites are soil-nesting termites, but how do they differ from typical soil-nesting termites?",
     "o":["Must permanently contact soil","Can survive fully脱离地面 without contacting soil",
          "Do not build tunnels","Do not swarm"]},
 28: {"q":"Which of the following is NOT a common sign of termite infestation?",
     "o":["Damaged wooden articles","Termite mud tubes","Hearing a 'beep' alarm sound","Large water stains on walls"]},
 29: {"q":"What are the main active ingredient and concentration of the Sentricon bait?",
     "o":["Fipronil 0.5%","Novaluron 0.5%","Cypermethrin 1%","Permethrin 0.5%"]},
 30: {"q":"How do termite members mainly communicate, guiding companions to suitable places?",
     "o":["Sound waves","Pheromone","Electric shock","Light"]},
},
"4": {
 0: {"s":"What is the definition of an accident?",
     "st":{"I":"An accident is one or a series of unplanned events causing injury, occupational disease, or damage to equipment or property, or environmental harm.",
           "II":"An accident is one or a series of unplanned events causing death, injury, occupational disease, or damage to equipment, or environmental harm.",
           "III":"An accident is one or a series of unplanned events causing death, injury, occupational disease, or damage to equipment or property, or environmental harm.",
           "IV":"An accident is a series of unplanned events causing death, injury, occupational disease, or damage to equipment or property, or environmental harm."},
     "o":["An accident is one or a series of unplanned events causing injury, occupational disease, or damage to equipment or property, or environmental harm.",
          "An accident is one or a series of unplanned events causing death, injury, occupational disease, or damage to equipment, or environmental harm.",
          "An accident is one or a series of unplanned events causing death, injury, occupational disease, or damage to equipment or property, or environmental harm.",
          "An accident is a series of unplanned events causing death, injury, occupational disease, or damage to equipment or property, or environmental harm."]},
 1: {"s":"Why does the company choose a full-body harness (also called parachute-style safety harness)?",
     "st":{"I":"Requirement of the pest control industry","II":"Lower price",
           "III":"Complies with ISO requirements","IV":"Complies with management company requirements"}},
 2: {"s":"When a fire occurs in a confined, poorly ventilated place, which extinguisher must we NOT choose?",
     "st":{"I":"Water extinguisher","II":"Carbon dioxide extinguisher","III":"Dry powder extinguisher","IV":"Foam extinguisher"},
     "o":["Water extinguisher","Carbon dioxide extinguisher","Dry powder extinguisher","Foam extinguisher"]},
 3: {"q":"What disease can be prevented by washing hands after handling chemicals?",
     "o":["Type A hepatitis","Forearm cramps","Dermatitis","Bronchitis"]},
 4: {"s":"When using a straight ladder, what ratio of base-to-height inclination must be maintained?",
     "st":{"I":"1:3","II":"1:4","III":"1:5","IV":"1:6"},
     "o":[None,None,None,None]},
 5: {"q":"At what height and above should appropriate work platforms be used?",
     "o":["1.5 m","2 m","2.5 m","3 m"]},
 6: {"s":"What is the most appropriate way to wear a safety harness?",
     "st":{"I":"Low hang, low use","II":"High hang, high use","III":"Low hang, high use","IV":"High hang, low use"},
     "o":["Low hang, low use","High hang, high use","Low hang, high use","High hang, low use"]},
 7: {"s":"During normal pest treatment, a client asks you to go out onto the bamboo scaffold to spray the exterior wall (directly through the window). The client communicates in Cantonese and says they have confirmed with the company an extra $1000. What would you do?",
     "st":{"I":"Return to the vehicle to get the safety harness, independent lifeline and fall arrester, then go onto the scaffold to finish the work.",
           "II":"Check whether the scaffold is stable, then go out to work.",
           "III":"Call the company to confirm whether an extra $1000 is acceptable; if yes, go out onto the scaffold to finish the work."}},
 8: {"s":"Which of the following is NOT a correct method during rodent treatment?",
     "st":{"I":"Must wear gloves and mask at work","II":"Must not directly touch rat bodies",
           "III":"Discard rat carcasses immediately","IV":"Wash hands thoroughly immediately after work",
           "V":"To protect employees' safety, what ordinance did the government enact and implement on 23 May 1997?",
           "VI":"Occupational Protection Health Ordinance","VII":"Occupational Accident Safety and Health Ordinance","VIII":"Occupational Protection Safety Ordinance"},
     "o":["Must wear gloves and mask at work","Must not directly touch rat bodies","Discard rat carcasses immediately","Wash hands thoroughly immediately after work"]},
 9: {"q":"Where should a smoke egg be lit when treating a location?",
     "o":["In the middle","At the farthest end","At the entrance/exit","Anywhere"]},
 10: {"s":"Which is NOT a potential hazard from incorrect chemical handling?",
     "st":{"I":"Cause a fire","II":"Cause injury to the lower back","III":"Splash corrosive or toxic liquid",
           "IV":"Cause injury, ulcers, poisoning or even death"},
     "o":["Cause a fire","Cause injury to the lower back","Splash corrosive or toxic liquid","Cause injury, ulcers, poisoning or even death"]},
 11: {"q":"What type of fire extinguisher is in the company vehicle?",
     "o":["Dry powder","Water","Carbon dioxide","Foam"]},
 12: {"q":"The main reason the company does not choose carbon dioxide extinguishers is:",
     "o":["ISO requirement","Higher price","It can cause asphyxiation to the user","It cannot isolate all fire sources"]},
 13: {"q":"Call the company to confirm whether an extra $1000 can be charged; if so, go out onto the scaffold to finish the work.",
     "o":[None,None,None,None]},
 14: {"s":"What must a chemical label possess?",
     "st":{"I":"Symbol of hazard classification","II":"Chemical name","III":"Indication of hazard situation",
           "IV":"Composition of the chemical","V":"Safety measures to be taken"}},
 15: {"q":"Try to list at least one potential danger of working at height, and the reason we cannot work on bamboo scaffolds.",
     "o":["Work platform or scaffold collapses causing the worker to fall",
          "Worker falls from work platform, scaffold or other elevated position",
          "Because insurance does not cover working on bamboo scaffolds"]},
 16: {"q":"In sudden high-risk situations per instructions from superiors or VIP clients, if we have no protective equipment, how should we handle it?",
     "o":["Solve the client's problem at all costs because we are professionals",
          "Where possible, handle as usual without protection because we are professionals",
          "Photograph the real situation and notify the company; only start work after all protection and equipment are in place",
          "The boss and client are always right; just do it without negotiation"]},
 17: {"s":"What is the correct way to use a turbo atomiser (Turbo) during treatment?",
     "st":{"I":"After observing, directly increase fan speed to cover the whole house with BKU.",
           "II":"First lay electric wires throughout the house to avoid repeated trips and reduce accidents and errors.",
           "III":"Think before acting; control wind force and chemical amount; do not let the machine hit client objects and decorations.",
           "IV":"Check the shoulder strap is secure; control wind force and chemical amount; spray unstable-wind objects."},
     "o":["After observing, directly increase fan speed to cover the whole house with BKU",
          "First lay electric wires throughout the house to avoid repeated trips and reduce accidents and errors",
          "Think before acting; control wind force and chemical amount; do not let the machine hit client objects and decorations"]},
 18: {"s":"According to company guidelines, why can we not work on bamboo scaffolds?",
     "st":{"I":"No formal training","II":"Insurance does not cover this work, so we get no protection",
           "III":"Extra fee required to do it","IV":"Because people fear heights; can try but not recommended"},
     "o":["No formal training","Insurance does not cover this work, so we get no protection","Extra fee required to do it","Because people fear heights; can try but not recommended"]},
 19: {"s":"What are the three basic elements of fire? Choose the most correct answer.",
     "st":{"I":"Fuel, air and heat source","II":"Rice straw, wooden sticks, dry places",
           "III":"Petrol, lighter, paper scraps","IV":"Accelerant, matches, ventilated place"},
     "o":["Fuel, air and heat source","Rice straw, wooden sticks, dry places","Petrol, lighter, paper scraps","Accelerant, matches, ventilated place"]},
 20: {"q":"What must a chemical label possess?",
     "o":[None,None,None,None]},
 21: {"q":"For unclear instructions from superiors or clients, high-risk work may arise; if we have no protective equipment, how should we handle it?",
     "o":["Solve the client's problem at all costs because we are professionals",
          "Where possible, handle as usual without protection because we are professionals",
          "Photograph the real situation and notify the company; only start work after all protection and equipment are in place",
          "The boss and client are always right; just do it without negotiation"]},
 22: {"s":"What must be noted when placing smoke generators (smoke eggs) or hot foggers during treatment?",
     "st":{"I":"Must wear an N95 mask","II":"Notify people on site and management, open windows",
           "III":"Start lighting the smoke generator (smoke egg) or hot fog from the farthest end",
           "IV":"Post a 'fumigation in progress' notice outside the door with the MSDS attached",
           "V":"Ensure no people or pets are present during fumigation"}},
 23: {"q":"What should be noted before and during hornet nest removal treatment? Find the correct answer.",
     "o":["I+II+II","II+III+IV","I+II+III","All of the above"]},
 24: {"q":"List precautions for termite liquid treatment (four or more). (4 points)",
     "o":["Think before acting; beware that the client's wooden items may be damaged and fall",
          "Check the drill's wiring is good to avoid electric leakage; hands must be dry when using the drill",
          "Mind the client's wiring layout to avoid damage causing shock; keep proper posture and back straight at work",
          "For concrete surfaces, mind various site dangers and obey safety rules",
          "Pay attention to the surroundings from multiple angles during work"]},
 25: {"q":"How is unsafe behaviour defined during work?",
     "o":["Wearing incomplete personal protective equipment","Using safe tools or machines",
          "Moving or placing objects and tools by safe methods","A colleague drinking water during lunch"]},
 26: {"s":"If a fire occurs near paper, flammable liquids and power sources, which extinguisher is more suitable?",
     "st":{"I":"Water extinguisher","II":"Carbon dioxide extinguisher","III":"Dry powder extinguisher","IV":"Foam extinguisher"},
     "o":["Water extinguisher","Carbon dioxide extinguisher","Dry powder extinguisher","Foam extinguisher"]},
},
"7": {
 0: {"s":"What types of rodents are commonly found in local homes?",
     "st":{"I":"Norway rat","II":"Black rat","III":"House mouse","IV":"Roof rat"}},
 1: {"q":"Black rat: also called roof rat, ship rat, black rat, yellow-breasted rat \u2014 where do they most appear?",
     "o":["In high floors, roof voids, pipes and gutters","In mid-low floors, roofs, empty rooms, pipes and gutters",
          "In low floors, roof voids, pipes and gutters","In mid-low-high floors, roof voids, pipes and gutters"]},
 2: {"s":"What are the behaviours and characteristics of rats?",
     "st":{"I":"Each group occupies a clear territory; adult rats attack all intruders",
           "II":"Live in groups or with their own kind","III":"Full of curiosity about novel objects",
           "IV":"Each group is dominated by a male"}},
 3: {"s":"What are the skills and habits of rats?",
     "st":{"I":"Excellent balance; easily land on both feet when jumping from height","II":"Good jumping ability, can jump up to 1 m high",
           "III":"Good at climbing and swimming","IV":"Range: rats 3-5 m, mice 50-100 m"}},
 4: {"s":"What are the sensory abilities of rats?",
     "st":{"I":"Keen sense of smell; can clearly remember the odour of poison bait","II":"Acute hearing; can hear ultrasound",
           "III":"Can distinguish colony members, markers of activity range and food source locations","IV":"Good vision, especially reactive to yellow and green objects"}},
 5: {"q":"Usually rats are not found during inspection, so we should look for the traces they leave.",
     "o":["Faeces, pathways, rat holes, rub marks, scars","Faeces, pipes, rat holes, rub marks, bite marks",
          "Faeces, pathways, rat holes, rub marks, bite marks","Faeces, odour, hair, rub marks, bite marks"]},
 6: {"q":"How far apart is it suitable to place bait stations on the perimeter?",
     "o":["5-10 m","5-10 feet","10-15 m","10-15 feet"]},
 7: {"q":"Approximately how long is the Norway rat body (head to tail)?",
     "o":["25-45 cm","30-45 cm","35-45 cm","40-45 cm"]},
 8: {"s":"When rat glue in a restaurant kitchen does not work fully, how should we follow up?",
     "st":{"I":"Try to block their entry points (ask client to block where found; technician handles reachable ones)",
           "II":"Cut off their food source","III":"Place rat traps","IV":"Place rat snap traps"}},
 9: {"q":"How many offspring can rats produce per year?",
     "o":["3-5 litters/year, 5-14 pups per litter","5-7 litters/year, 5-14 pups per litter",
          "6-9 litters/year, 5-14 pups per litter","5-8 litters/year, 5-14 pups per litter"]},
 10: {"q":"Why is the rat bait station warning printed in red? Find the correct answer.",
     "o":["They are especially sensitive to and attracted by red","Aesthetic; matches the bait station design",
          "Red represents danger, warns others not to touch","They are colour-blind"]},
 11: {"q":"A resident hired a pest company to eradicate rats indoors; a child was taken to hospital after ingesting rat poison. Which material do you estimate the company failed to use?",
     "o":["Mouse trap","Rat glue","Rat bait station","Rat cage"]},
 12: {"q":"A property with mosquito nets installed and air-conditioning left on, after many fly treatments still has fly problems but also many dead flies \u2014 which is the reason?",
     "o":["Insecticide diluted too weakly","Cold fog (Turbo) sprayed unevenly","Insecticide efficacy has degraded",
          "Breeding source still not addressed"]},
 13: {"s":"What steps should be noted when replacing rat bait? Find the correct answer.",
     "st":{"I":"Replace damaged and unclean warning labels","II":"No need to replace old mouldy bait to avoid waste and support the environment",
           "III":"Set bait stations where rats are active, foraging and nesting","IV":"Ensure the bait station is intact and bait stored securely"}},
 14: {"s":"Which method is the basic control for rodent infestation?",
     "st":{"I":"Cut off rat food source","II":"Place rat traps and/or snap traps","III":"Set up bait stations","IV":"Install ultrasound"},
     "o":["Cut off rat food source","Place rat traps and/or snap traps","Set up bait stations","Install ultrasound"]},
 15: {"s":"What are the characteristics of fruit flies?",
     "st":{"I":"About 3-4 mm long","II":"Black body","III":"Usually red eyes","IV":"Eat fresh fruit, vinegar, wine"}},
 16: {"s":"What are the characteristics of moth flies?",
     "st":{"I":"Hairy and slender","II":"Yellow, grey-brown to deep black","III":"Triangular wings","IV":"Covered with fine hairs"}},
 17: {"s":"Where should we focus inspection for house fly nuisance? Find the correct answer.",
     "st":{"I":"They do not often breed indoors, so likely come from outside","II":"Inspect external structures - surroundings",
           "III":"Look for garbage areas or decaying organic matter","IV":"Inspect changing rooms, storage rooms, etc."}},
 18: {"q":"Norway rat, also called brown rat, common rat, Norway rat \u2014 where do they most appear?",
     "o":["In high floors, roof voids, pipes and gutters","In sewers, low floors, tree holes and grass",
          "In sewers, low floors, burrows","Appear anywhere"]},
 19: {"s":"What should be noted for personal hygiene when handling rat carcasses? Find the correct answer.",
     "st":{"I":"All clothing and items contaminated by dead rats must be thoroughly disinfected",
           "II":"Gloves need cleaning, disinfection or disposal","III":"Masks need cleaning, disinfection or disposal",
           "IV":"Hold breath tightly to avoid inhaling bacteria"}},
 20: {"q":"What is the main ingredient of our commonly used Contrac Blox (blue wax block) rodent bait?",
     "o":[None,None,None,None]},
 21: {"s":"What are the characteristics of the House Fly?",
     "st":{"I":"About 4-7.5 mm long","II":"Dark grey body","III":"Usually red eyes","IV":"Four vertical stripes on abdomen"}},
 22: {"q":"To which order do rats belong?",
     "o":["Carnivora","Rodentia","Chiroptera","Artiodactyla"]},
 23: {"q":"Identifying Norway rat vs roof rat, which description is correct?",
     "o":["Norway rat tail longer than head+body; roof rat tail shorter than head+body",
          "Norway rat tail shorter than head+body, heavy body; roof rat tail longer than head+body, slender body",
          "Both have the same tail length","Norway rat has large ears; roof rat has small ears"]},
 24: {"q":"Which of the following about rat activity habits is correct?",
     "o":["Active by day, resting at night","Live alone, no territoriality",
          "Active at night, resting by day, and establish protective routes along walls or covered areas","No wariness of novel objects"]},
 25: {"q":"Which of the following is NOT a common rodent-borne disease?",
     "o":["Plague","Hantavirus","Hepatitis E","Dengue fever"]},
 26: {"q":"Active rat burrows usually have which characteristics?",
     "o":["Only 1 exit, rough hole","2-3 exits, 50 cm deep, smooth hole with faeces/food debris",
          "Definitely in the ceiling","Spider webs at the hole"]},
 27: {"q":"What class do our commonly used rodenticides mainly belong to?",
     "o":["Neurotoxins","Anticoagulants","Muscle relaxants","Emetics"]},
 28: {"q":"The antidote for anticoagulant poisoning is?",
     "o":["Vitamin C","Vitamin K1","Salt","Glucose"]},
 29: {"q":"Why do we not use toxic bait indoors?",
     "o":["Indoor bait is ineffective","Human and pet safety; if a rat dies in an inaccessible place it causes serious hygiene problems",
          "Indoor use is legally prohibited","Bait fears being eaten by koalas"]},
 30: {"q":"Which of the following about the rodent bait station is incorrect?",
     "o":["Designed to allow only rodents to enter and feed","Prevents non-target animals and humans from tampering or entering",
          "Can have warning and inspection labels attached","Can place cat/dog food inside to attract rats"]},
 31: {"q":"The metamorphosis of flies belongs to?",
     "o":["Incomplete metamorphosis","Complete metamorphosis","No metamorphosis","Hemimetabolism"]},
 32: {"q":"Which is NOT a problem caused by flies?",
     "o":["Nuisance","Disease transmission","Lay eggs and grow in food sources","Mainly bite and suck blood"]},
},
"8": {
 0: {"q":"Which cockroach especially likes starchy food?",
     "o":["Long-tailed cockroach (brown-banded)","Long-tailed cockroach (brown-banded) and German cockroach",
          "Long-tailed cockroach (brown-banded) and Australian cockroach","Long-tailed cockroach (brown-banded) and American cockroach"]},
 1: {"s":"What are the life cycle, behaviour and habits of cockroaches?",
     "st":{"I":"Extremely resilient","II":"Can grow with only a small amount of food",
           "III":"Can survive long periods without food and water","IV":"More active in colder places"}},
 2: {"s":"What are the main control methods for German cockroaches?",
     "st":{"I":"Inspection and advice to client","II":"Spray the whole house vigorously and apply Top Bait Gel",
           "III":"Use pheromone cockroach traps","IV":"Reduce spraying because they concentrate in active spots and develop resistance in the next generation"}},
 3: {"q":"Find the correct answer for the life cycle, behaviour and habits of German cockroaches.",
     "o":["I+II","I+II+III","I+II+IV","All of the above"]},
 4: {"q":"Find the correct answer for the life cycle, behaviour and habits of American cockroaches.",
     "o":["They only fly when temperature is above 30 C","Female lives 14-15 months",
          "Male lives shorter","All of the above"]},
 5: {"q":"How many egg capsules can a female German cockroach produce in her lifetime, and how many offspring?",
     "o":["2-6 capsules, 300 offspring","6-10 capsules, 250 offspring",
          "4-8 capsules, 300 offspring","4-8 capsules, 400 offspring"]},
 6: {"q":"Find the correct answer for the life cycle, behaviour and habits of American cockroaches.",
     "o":["I+II+III","II+III+IV","All of the above","None of the above"]},
 7: {"q":"How is the brown-banded cockroach (commonly called 'coffee-band') treated?",
     "o":["First treat targeted and apply chemical in all hidden positions, do cold-fog for the whole house, thinly spray inside cabinets if needed.",
          "Apply gel vigorously, no spraying, use steam (high-temperature treatment) and vacuum cleaner.",
          "Do everything, heavy fog, set pheromone in the house, add chemical to gel.",
          "Call the company for advice and request extra time; next treatment delayed 30 minutes."]},
 8: {"q":"Find the correct answer for the life cycle, behaviour and habits of German cockroaches.",
     "o":["Colour - light to medium brown, length 1.5-2 cm, winged but cannot fly, life up to 12 months.",
          "Colour - light to dark brown, length 2.5-3 cm, good at flying, life up to 12 months.",
          "Colour - light to medium brown, length 2.5-3 cm, cannot fly, life up to 12 months.",
          "Colour - light to dark brown, length 1.5-2 cm, good at flying, life up to 12 months."]},
 9: {"q":"Which German cockroaches are especially attracted by pheromones?",
     "o":["Male","Female","Immature","All"]},
 10: {"q":"From which two features can German cockroaches and brown-banded cockroaches be distinguished?",
     "o":["Two dark obvious stripes on the head and body size","Three dark obvious stripes on the head and body size",
          "Two dark obvious stripes on the head and whether they scatter","Two dark obvious stripes on the head and number of legs"]},
 11: {"q":"Which two methods are suitable for controlling all species of cockroaches?",
     "o":["Spray insecticide and use bait","Use bait and cockroach traps (houses)",
          "Inspection, advice to client and vacuuming","Spray insecticide and use cockroach traps (houses)"]},
 12: {"q":"How to identify silverfish? Find the correct answer.",
     "o":["12-25 mm, coffee to silver-grey, carrot-shaped","12-25 mm, dark coffee to iron-grey, carrot-shaped",
          "15-35 mm, coffee to silver-grey, carrot-shaped","15-35 mm, dark coffee to iron-grey, carrot-shaped"]},
 13: {"q":"Which two insects mainly feed on mould?",
     "o":["I + II","I + III","II + IV","III + IV"]},
 14: {"q":"How to control booklice (generally)?",
     "o":["Remove damp conditions, lower relative humidity to 50%, keep food dry and vacuum if needed.",
          "Remove dry conditions, raise relative humidity to 90%, advise client to keep home clean.",
          "Remove dry conditions, raise relative humidity to 90%, advise client to keep home clean and spray BK-Mite vigorously.",
          "Remove damp conditions, lower relative humidity to 50%-70%, keep food dry and vacuum if needed."]},
 15: {"q":"A bug about the length of a fingernail, grey, slim, with three tails at the rear \u2014 which is it?",
     "o":["Springtail","Earwig","Silverfish","German cockroach"]},
 16: {"q":"When German cockroaches are found, we generally apply gel and insect monitoring stations. Where should we apply gel according to our practice?",
     "o":["Anywhere, since cockroaches move randomly.","In conspicuous places so the client knows we followed up.",
          "In hidden cabinet positions, hinges, corners or inside skirting boards and other hidden hiding places.","Sewers, drains or window sills."]},
 17: {"q":"Spraying insecticide is a major part of American cockroach treatment, but which position should be minimally sprayed directly?",
     "o":["Unrelated to insecticide safety","Master bedroom","Baby room","Dog house","Inside sewer"]},
 18: {"q":"Under what circumstances do earwigs generally enter the home environment?",
     "o":["Brought in via plants or flowers, entering through gaps or cracks.","Sewers, drains or damp places.",
          "Any externally brought item, e.g. delivery logistics items.","All of the above"]},
 19: {"q":"A client with a small dog has a 4-time booklice contract, two treatments done (weekly, on time); problem much improved. For the third treatment, which correct method helps the client solve the booklice problem?",
     "o":["Spray walls with Handpump and BKL","Spray Turbo with BKL","Persuade her to renew contract",
          "Educate her to control humidity and reduce mould growth"]},
},
}

# ----------------------------------------------------------------------------
# Build converted questions
# ----------------------------------------------------------------------------
def norm_combo_option(o):
    s = o.strip()
    s = re.sub(r"^[A-Da-d][\)\.、]?\s*", "", s)
    s = re.sub(r"\s*\+\s*", "+", s)
    s = re.sub(r"([IVX]+)\s*\.\s*", r"\1", s)
    if s == "以上皆是":
        return "以上皆是 (All of the above)"
    if s == "以上皆不是":
        return "以上皆不是 (None of the above)"
    return s

def wrap_plain_option(o, en):
    if en is None:
        return o
    # if already bilingual (Chinese + English in parentheses), keep
    if "(" in o and any(c.isascii() and c.isalpha() for c in o):
        return o
    return f"{o} ({en})"

def build_question(t, idx, q_obj):
    if t == "2" and idx <= 4:
        return q_obj["question"], q_obj["options"], "kept"

    is_combo = is_combo_options(q_obj["options"])
    qtext = q_obj["question"]
    entry = T.get(t, {}).get(idx)

    # special: 4-20 borrows statements from 4-14
    if t == "4" and idx == 20:
        stem_cn, _ = parse_statements(qtext)
        stmts_cn = parse_statements(data["4"]["mc"][14]["question"])[1]
        stem_en = entry["q"]
        stmts_en = T["4"][14]["st"]
        lines = [stem_cn, stem_en]
        for roman, cn in stmts_cn:
            lines.append(""); lines.append(f"{roman}. {cn}"); lines.append(f"{roman}. {stmts_en[roman]}")
        return "\n".join(lines), [norm_combo_option(o) for o in q_obj["options"]], "warn_repair"

    stem_cn, stmts_cn = parse_statements(qtext)
    has_stmt = len(stmts_cn) > 0

    if has_stmt:
        lines = [stem_cn, entry["s"]]
        for roman, cn in stmts_cn:
            lines.append(""); lines.append(f"{roman}. {cn}"); lines.append(f"{roman}. {entry['st'][roman]}")
        new_q = "\n".join(lines)
        if is_combo:
            return new_q, [norm_combo_option(o) for o in q_obj["options"]], "ai"
        otrans = entry.get("o", [None]*len(q_obj["options"]))
        new_opts = [wrap_plain_option(o, otrans[i]) for i, o in enumerate(q_obj["options"])]
        return new_q, new_opts, "ai"
    else:
        if "q" not in entry:
            print("MISSING q for", t, idx, "keys=", list(entry.keys()), "qtext=", repr(qtext[:60]))
            return qtext, q_obj["options"], "ai"
        new_q = qtext + "\n" + entry["q"]
        if is_combo:
            return new_q, [norm_combo_option(o) for o in q_obj["options"]], "ai"
        otrans = entry.get("o", [None]*len(q_obj["options"]))
        new_opts = [wrap_plain_option(o, otrans[i]) for i, o in enumerate(q_obj["options"])]
        return new_q, new_opts, "ai"

def detect_mismatch(t, idx, q_obj):
    if not is_combo_options(q_obj["options"]):
        return None
    if t == "2" and idx <= 4:
        return None
    if t == "4" and idx in (13, 20):
        return None
    _, stmts_cn = parse_statements(q_obj["question"])
    n_stmt = len(stmts_cn)
    mx = max_roman_in_options(q_obj["options"])
    if mx > n_stmt:
        rn = {4:"IV",5:"V",8:"VIII"}.get(mx, str(mx))
        return f"statement 數量不匹配 (options 用至 {rn}, 題目只有 {n_stmt} 個 statement)"
    return None

# ----------------------------------------------------------------------------
# Produce converted structure + tags
# ----------------------------------------------------------------------------
out = {}
tags = {}   # (t,idx) -> marker text
warnings_docx = []

for t in data:
    topic = data[t]
    out[t] = {"name": topic["name"], "mc": [], "essay": copy.deepcopy(topic.get("essay", []))}
    for idx, q_obj in enumerate(topic["mc"]):
        if t == "2" and idx <= 4:
            out[t]["mc"].append(copy.deepcopy(q_obj))
            tags[(t, idx)] = "［你已改·保留］"
            continue
        new_q, new_opts, marker = build_question(t, idx, q_obj)
        out[t]["mc"].append({"question": new_q, "options": new_opts, "correct": q_obj["correct"]})
        if marker == "warn_repair":
            tags[(t, idx)] = "［⚠️ 修復·請確認］ (statement 由 4-14 補入)"
        elif marker == "kept":
            tags[(t, idx)] = "［你已改·保留］"
        else:
            tags[(t, idx)] = "［AI 轉換］"

# warnings for mismatch + special
for t in data:
    for idx, q_obj in enumerate(data[t]["mc"]):
        if t == "2" and idx <= 4:
            continue
        if t == "4" and idx == 13:
            warnings_docx.append((t, idx, "［⚠️ 原題目文本不完整·請用戶自行修正］ (題幹僅一句，無 statement)"))
            continue
        if t == "4" and idx == 20:
            warnings_docx.append((t, idx, "［⚠️ 修復·請確認］ (statement 由 4-14 補入)"))
            continue
        m = detect_mismatch(t, idx, q_obj)
        if m:
            warnings_docx.append((t, idx, "［⚠️ " + m + "·請確認］"))

# ----------------------------------------------------------------------------
# Write JSON (UTF-8, no BOM)
# ----------------------------------------------------------------------------
with open(OUT_JSON, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=2)
    f.write("\n")

# ----------------------------------------------------------------------------
# Self-verification
# ----------------------------------------------------------------------------
total = 0
per_topic = {}
combo_count = 0
plain_count = 0
kept_count = 0
for t in data:
    n = len(out[t]["mc"])
    per_topic[t] = n
    total += n
    for idx, q in enumerate(out[t]["mc"]):
        if t == "2" and idx <= 4:
            kept_count += 1
            continue
        if is_combo_options(data[t]["mc"][idx]["options"]):
            combo_count += 1
        else:
            plain_count += 1

# options count + correct integrity
opt_ok = True
correct_ok = True
for t in data:
    for idx, q in enumerate(data[t]["mc"]):
        if len(out[t]["mc"][idx]["options"]) != len(q["options"]):
            opt_ok = False
            print("OPT MISMATCH", t, idx)
        if out[t]["mc"][idx]["correct"] != q["correct"]:
            correct_ok = False
            print("CORRECT MISMATCH", t, idx)

# preserved deep-equal
preserved_ok = True
for idx in range(5):
    a = json.dumps(data["2"]["mc"][idx], ensure_ascii=False, sort_keys=True)
    b = json.dumps(out["2"]["mc"][idx], ensure_ascii=False, sort_keys=True)
    if a != b:
        preserved_ok = False
        print("PRESERVED MISMATCH", idx)

print("=== VERIFICATION ===")
print("Total MC:", total)
print("Per topic:", per_topic)
print("Kept(preserved):", kept_count, "Combo(converted):", combo_count, "Plain(converted):", plain_count)
print("Options count consistent:", opt_ok)
print("Correct unchanged:", correct_ok)
print("Topic2 idx0-4 byte-identical(content):", preserved_ok)
print("Warnings (⚠️):", len(warnings_docx))
for w in warnings_docx:
    print("   ", w[0]+"-"+str(w[1]), w[2])

# ----------------------------------------------------------------------------
# DOCX preview
# ----------------------------------------------------------------------------
os.makedirs(os.path.dirname(OUT_DOCX), exist_ok=True)
from docx import Document

doc = Document()
doc.add_heading("BIOCYCLE 題庫中英對照 — 最終預覽（MC）", level=0)
doc.add_paragraph("Final Bilingual MC Preview — Pending Confirmation")
doc.add_paragraph("本預覽與將上線之最終版本完全一致（mc 部分）。essay 不在本文件內。")
doc.add_paragraph("")

# warnings block at top
doc.add_heading("⚠️ 需要重點檢查的題目（請用戶確認）", level=1)
if warnings_docx:
    for t, idx, reason in warnings_docx:
        doc.add_paragraph(f"{t}-{idx}: {reason}")
else:
    doc.add_paragraph("（無）")

for t in data:
    doc.add_heading(out[t]["name"], level=1)
    for idx, q in enumerate(out[t]["mc"]):
        qlines = q["question"].split("\n")
        p = doc.add_paragraph()
        p.add_run(f"{idx+1}. ")
        for j, line in enumerate(qlines):
            if j > 0:
                p.add_run().add_break()
            p.add_run(line)
        for i, opt in enumerate(q["options"]):
            disp = re.sub(r"^[A-Da-d][\)\.]\s*", "", opt)
            doc.add_paragraph(f"{chr(65+i)}. {disp}")
        rc = doc.add_paragraph()
        run = rc.add_run(f"正確答案：{chr(65+q['correct'])}")
        run.bold = True
        doc.add_paragraph(tags[(t, idx)])
        doc.add_paragraph("")

doc.save(OUT_DOCX)
print("DOCX written:", OUT_DOCX)
print("DONE")
